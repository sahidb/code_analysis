import os
import json
import subprocess
import argparse
from datetime import datetime
from docx import Document
from docx.shared import RGBColor

def run_eslint(file_path):
    """
    Run ESLint on a specific file and capture the JSON output.
    """
    try:
        result = subprocess.run(
            ["eslint", "--format", "json", file_path],
            capture_output=True,
            text=True,
            check=True
        )
        return json.loads(result.stdout)
    except subprocess.CalledProcessError as e:
        print(f"Error running ESLint on {file_path}: {e}")
        return []

def analyze_project_folder(folder_path, exclude_folders=None):
    if exclude_folders is None:
        exclude_folders = ['node_modules']  # Default folder to exclude

    all_metrics = []
    
    for root, dirs, files in os.walk(folder_path):
        # Exclude specified folders
        dirs[:] = [d for d in dirs if d not in exclude_folders]
        
        for file in files:
            if file.endswith((".js", ".jsx", ".ts", ".tsx")):  # Analyze JS, JSX, TS, TSX files
                file_path = os.path.join(root, file)
                print(f"Analyzing: {file_path}")
                eslint_results = run_eslint(file_path)
                all_metrics.extend(eslint_results)  # Combine results for each file
    
    return all_metrics

def generate_word_report(analysis_results, output_path):
    if not analysis_results:
        print("No data to generate a Word report.")
        return
    
    doc = Document()
    doc.add_heading('React Code Analysis Report', 0)
    doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    explanation = (
        "This report provides an analysis of the React code within the specified folder. "
        "The analysis includes linting results, which identify coding style issues, errors, and warnings."
    )
    doc.add_paragraph(explanation)

    doc.add_heading('ESLint Analysis Results', level=1)
    
    for result in analysis_results:
        doc.add_heading(f"File: {result['filePath']}", level=2)
        
        for message in result['messages']:
            row_text = f"Line {message['line']}: [{message['severity']}] {message['message']} (rule: {message['ruleId']})"
            para = doc.add_paragraph(row_text)
            
            if message['severity'] == 2:
                para.font.color.rgb = RGBColor(255, 0, 0)  # Red for errors
            else:
                para.font.color.rgb = RGBColor(255, 165, 0)  # Orange for warnings
    
    doc.save(output_path)
    print(f"Word report generated: {output_path}")

def generate_html_report(analysis_results, output_path):
    if not analysis_results:
        print("No data to generate an HTML report.")
        return
    
    html_content = """
    <html>
    <head>
        <title>React Code Analysis Report</title>
        <style>
            body { font-family: Arial, sans-serif; margin: 20px; background-color: #f9f9f9; }
            .container { max-width: 1200px; margin: auto; background-color: #fff; padding: 20px; box-shadow: 0px 0px 10px rgba(0, 0, 0, 0.1); }
            h1, h2, h3 { color: #333; }
            table { width: 100%; border-collapse: collapse; margin-bottom: 20px; }
            th, td { padding: 10px; border: 1px solid #dddddd; text-align: left; }
            th { background-color: #f2f2f2; font-weight: bold; }
            .error { background-color: #f8d7da; color: #721c24; }
            .warning { background-color: #fff3cd; color: #856404; }
        </style>
    </head>
    <body>
        <div class="container">
            <h1>React Code Analysis Report</h1>
            <p>Generated on: {date}</p>
            <h2>ESLint Analysis Results</h2>
    """.format(date=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    
    for result in analysis_results:
        html_content += f"<h3>File: {result['filePath']}</h3><table><tr><th>Line</th><th>Severity</th><th>Message</th><th>Rule</th></tr>"
        
        for message in result['messages']:
            severity_class = "error" if message['severity'] == 2 else "warning"
            html_content += f"""
                <tr class="{severity_class}">
                    <td>{message['line']}</td>
                    <td>{'Error' if message['severity'] == 2 else 'Warning'}</td>
                    <td>{message['message']}</td>
                    <td>{message['ruleId']}</td>
                </tr>
            """
        
        html_content += "</table>"
    
    html_content += """
        </div>
    </body>
    </html>
    """
    
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print(f"HTML report generated: {output_path}")

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Generate React code analysis reports in HTML or Word format")
    parser.add_argument('folder_path', type=str, help="Path to the React project src folder")
    parser.add_argument('--format', choices=['html', 'word'], default='html', help="The report format to generate (html or word)")
    parser.add_argument('--output', type=str, required=True, help="The output path for the report")
    args = parser.parse_args()
    
    folder_path = args.folder_path
    report_format = args.format
    output_path = args.output

    # Analyze the React project folder
    analysis_results = analyze_project_folder(folder_path, exclude_folders=['node_modules'])

    # Generate the report based on the selected format
    if report_format == 'html':
        generate_html_report(analysis_results, output_path)
    else:
        generate_word_report(analysis_results, output_path)
